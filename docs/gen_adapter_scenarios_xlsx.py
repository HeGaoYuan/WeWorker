#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
导出本地所有 adapter 的支持场景到 Excel，用于统计每个 adapter×场景各阶段耗时。

场景提取逻辑对齐 backend-ts/src/site-scenarios.ts：
  - manifest.json 的 `scenarios` 字段 = registry 源
  - 无 scenarios 时，按 description 短语 fallback 合成（同 synthesizeFallback）

阶段耗时列（LLM思考 / 工具执行 / 总耗时 / 执行步数）对应现有 timing 埋点
（works-store.ts 的 {totalMs, llmMs, execMs, steps}），当前留空待填。

用法:
  python gen_adapter_scenarios_xlsx.py [--adapters-dir DIR] [--out FILE]
"""
import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ---------- 场景解析（复刻 site-scenarios.ts）----------

def parse_description_phrases(description: str) -> list[str]:
    if not description:
        return []
    if "：" in description:
        tail = description.split("：", 1)[1]
    elif ":" in description:
        tail = description.split(":", 1)[1]
    else:
        tail = description
    out = []
    for part in re.split(r"[，,、]", tail):
        part = part.strip()
        if 2 <= len(part) <= 32:
            out.append(part)
    return out


def synthesize_fallback(name: str, description: str) -> list[dict]:
    defs, seen = [], set()
    for phrase in parse_description_phrases(description):
        if phrase in seen:
            continue
        seen.add(phrase)
        defs.append({
            "id": f"desc-{phrase}",
            "title": phrase,
            "summary": f"在{name}完成「{phrase}」",
            "steps": [],
            "tool_names": [],
            "prompts": [],
            "source": "fallback",
        })
    if not defs:
        defs.append({
            "id": "default",
            "title": "通用操作",
            "summary": description or "",
            "steps": [],
            "tool_names": [],
            "prompts": [],
            "source": "fallback",
        })
    return defs


def adapter_name(manifest: dict, adapter_id: str) -> str:
    return (
        manifest.get("i18n", {}).get("zh", {}).get("name")
        or manifest.get("name")
        or adapter_id
    )


def get_scenarios(manifest: dict, name: str) -> list[dict]:
    raw = manifest.get("scenarios")
    if isinstance(raw, list) and raw:
        out = []
        for s in raw:
            if not isinstance(s, dict):
                continue
            out.append({
                "id": s.get("id") or s.get("title") or "",
                "title": s.get("title") or "",
                "summary": s.get("summary") or "",
                "steps": s.get("steps") or [],
                "tool_names": s.get("tool_names") or [],
                "prompts": s.get("prompts") or [],
                "source": "registry",
            })
        if out:
            return out
    return synthesize_fallback(name, manifest.get("description", ""))


# ---------- 数据收集 ----------

def collect(adapters_dir: Path) -> list[dict]:
    rows = []
    for d in sorted(adapters_dir.iterdir()):
        manifest_path = d / "manifest.json"
        if not manifest_path.is_file():
            continue
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        except Exception as exc:  # noqa: BLE001 - skip malformed adapter (同 adapter-parser)
            print(f"skip {d.name}: {exc}", file=sys.stderr)
            continue

        adapter_id = manifest.get("id") or d.name
        name = adapter_name(manifest, adapter_id)
        requires_auth = "是" if manifest.get("requires_auth") else "否"
        auth_tool = manifest.get("auth_check_tool") or ""
        match = (manifest.get("match") or [adapter_id])[0]

        for sc in get_scenarios(manifest, name):
            prompts = sc["prompts"]
            first_prompt = ""
            if isinstance(prompts, list) and prompts:
                first = prompts[0]
                text = first.get("text", "") if isinstance(first, dict) else str(first)
                first_prompt = (text[:80] + "…") if len(text) > 80 else text
            rows.append({
                "adapter_id": adapter_id,
                "adapter_name": name,
                "requires_auth": requires_auth,
                "auth_check_tool": auth_tool,
                "match": match,
                "scenario_id": sc["id"],
                "scenario_title": sc["title"],
                "summary": sc["summary"],
                "source": sc["source"],
                "steps_n": len(sc["steps"]),
                "tools_n": len(sc["tool_names"]),
                "prompts_n": len(prompts),
                "first_prompt": first_prompt,
            })
    return rows


# ---------- Excel 写出 ----------

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TO_FILL_FILL = PatternFill("solid", fgColor="FFF2CC")  # 待填列浅黄底
BAND_FILL = PatternFill("solid", fgColor="EDEDED")     # adapter 交替带
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [
    ("序号", 6),
    ("Adapter ID", 22),
    ("Adapter名称", 16),
    ("需登录", 8),
    ("登录检测工具", 28),
    ("站点(match)", 22),
    ("场景ID", 20),
    ("场景标题", 18),
    ("场景摘要", 46),
    ("场景来源", 10),
    ("场景步骤数(定义)", 14),
    ("涉及工具数", 10),
    ("Prompts数", 10),
    ("首个Prompt预览", 50),
    ("LLM思考(ms)", 13),     # 待填 → thinkMs / llmMs
    ("工具执行(ms)", 13),    # 待填 → execMs
    ("总耗时(ms)", 12),      # 待填 → totalMs
    ("执行步数(埋点)", 13),  # 待填 → steps
    ("备注", 20),
]
TO_FILL_COL_INDEXS = {15, 16, 17, 18}  # 1-based 列号，待填耗时列


def build_main_sheet(ws, rows: list[dict]) -> None:
    # 表头
    headers = [title for title, _ in COLUMNS]
    ws.append(headers)
    for col_idx in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMNS[col_idx - 1][1]
    ws.row_dimensions[1].height = 30

    # 数据
    band = False
    prev_adapter = None
    for i, r in enumerate(rows, start=1):
        if r["adapter_id"] != prev_adapter:
            band = not band
            prev_adapter = r["adapter_id"]
        values = [
            i,
            r["adapter_id"],
            r["adapter_name"],
            r["requires_auth"],
            r["auth_check_tool"],
            r["match"],
            r["scenario_id"],
            r["scenario_title"],
            r["summary"],
            r["source"],
            r["steps_n"],
            r["tools_n"],
            r["prompts_n"],
            r["first_prompt"],
            None,  # LLM思考(ms)
            None,  # 工具执行(ms)
            None,  # 总耗时(ms)
            None,  # 执行步数
            None,  # 备注
        ]
        ws.append(values)
        row_idx = i + 1
        for col_idx in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = BORDER
            wrap = COLUMNS[col_idx - 1][0] in ("场景摘要", "首个Prompt预览", "登录检测工具", "备注")
            cell.alignment = Alignment(
                horizontal="center" if col_idx in (1, 4, 10, 11, 12, 13) else "left",
                vertical="center",
                wrap_text=wrap,
            )
            if col_idx in TO_FILL_COL_INDEXS:
                cell.fill = TO_FILL_FILL
            elif band:
                cell.fill = BAND_FILL

    ws.freeze_panes = "D2"  # 冻结首行 + 序号/Adapter ID/名称
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"


def build_readme_sheet(ws, rows: list[dict], adapters_dir: Path, out_path: Path) -> None:
    adapter_count = len({r["adapter_id"] for r in rows})
    registry_n = sum(1 for r in rows if r["source"] == "registry")
    fallback_n = sum(1 for r in rows if r["source"] == "fallback")

    lines = [
        ("Adapter×场景 耗时统计表 — 说明", True),
        ("", False),
        (f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", False),
        (f"数据来源: {adapters_dir}", False),
        (f"输出文件: {out_path}", False),
        ("", False),
        (f"Adapter 总数: {adapter_count}", False),
        (f"场景总行数: {len(rows)}", False),
        (f"  其中 registry 源(manifest.scenarios): {registry_n}", False),
        (f"  其中 fallback 源(按 description 合成): {fallback_n}", False),
        ("", False),
        ("列含义:", True),
        ("  序号 / Adapter ID / Adapter名称 / 需登录 / 登录检测工具 / 站点(match) — adapter 元信息", False),
        ("  场景ID / 场景标题 / 场景摘要 / 场景来源 — 场景信息(来源=registry 表示 manifest 显式定义;=fallback 表示按 description 合成)", False),
        ("  场景步骤数(定义) — manifest 里 scenario.steps 的长度(场景定义的步骤数，非实测)", False),
        ("  涉及工具数 / Prompts数 — scenario.tool_names / scenario.prompts 的长度", False),
        ("  首个Prompt预览 — 场景第一个 prompt 文本(截断 80 字)", False),
        ("", False),
        ("阶段耗时列(浅黄底，待填):", True),
        ("  LLM思考(ms)   ← 对应埋点 thinkMs/llmMs  (LLM 推理耗时)", False),
        ("  工具执行(ms)   ← 对应埋点 execMs        (工具调用执行耗时)", False),
        ("  总耗时(ms)    ← 对应埋点 totalMs       (任务整体耗时)", False),
        ("  执行步数(埋点) ← 对应埋点 steps         (agent loop 步数)", False),
        ("  说明: 现有 works.db 只存任务级 timing 且不绑定 adapter/场景，", False),
        ("        故此处留空；跑任务后手工填入实测值即可用于横向统计对比。", False),
        ("  统计方式: 填完后可直接对耗时列做 SUM/AVERAGE/透视，空单元格会被自动忽略。", False),
        ("", False),
        ("重跑更新: python docs/gen_adapter_scenarios_xlsx.py", False),
    ]
    ws.column_dimensions["A"].width = 100
    for idx, (text, bold) in enumerate(lines, start=1):
        cell = ws.cell(row=idx, column=1, value=text)
        cell.font = Font(bold=bold, size=12 if bold and idx == 1 else 11)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def main() -> int:
    repo_root = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(description="导出 adapter×场景到 xlsx")
    parser.add_argument("--adapters-dir", default=str(repo_root / "extension" / "adapters"))
    parser.add_argument("--out", default=str(Path(__file__).resolve().parent / "adapter_scenarios_timing.xlsx"))
    args = parser.parse_args()

    adapters_dir = Path(args.adapters_dir)
    out_path = Path(args.out)
    if not adapters_dir.is_dir():
        print(f"adapters dir not found: {adapters_dir}", file=sys.stderr)
        return 1

    rows = collect(adapters_dir)
    if not rows:
        print("no adapter scenarios collected", file=sys.stderr)
        return 1

    wb = Workbook()
    main_ws = wb.active
    main_ws.title = "场景耗时统计"
    build_main_sheet(main_ws, rows)
    readme_ws = wb.create_sheet("说明")
    build_readme_sheet(readme_ws, rows, adapters_dir, out_path)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    adapter_count = len({r["adapter_id"] for r in rows})
    print(f"OK -> {out_path}")
    print(f"  adapters: {adapter_count}, scenario rows: {len(rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
